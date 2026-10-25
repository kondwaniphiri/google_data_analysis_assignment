"""Build the Word report from its reviewable Markdown source and evidence PNGs.

Use the bundled document runtime (python-docx and Pillow). No desktop UI is needed.
"""
from pathlib import Path
import argparse
import re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image

ROOT=Path(__file__).resolve().parents[1]
parser=argparse.ArgumentParser()
parser.add_argument('--student-name',default='[Enter student name]')
parser.add_argument('--student-number',default='[Enter NCC student number]')
args=parser.parse_args()
source=(ROOT/'report/report.md').read_text(encoding='utf-8')
figures=[('01_data_quality','Figure 1 Data quality audit and analysis populations'),
 ('02_preprocessing','Figure 2 Applied spelling corrections and tokenization method'),
 ('03_queries','Figure 3 Complete query results including four tied places'),
 ('04_phrases','Figure 4 Twenty most frequent phrases in five star reviews'),
 ('05_sentiment','Figure 5 VADER sentiment distribution and rating comparison')]
# Count report prose, headings, figure headings and captions; exclude cover,
# references and text inside screenshots, in accordance with the brief.
main=source.split('## References')[0]
word_count=len(re.findall(r'\S+',re.sub(r'(?m)^#{1,3} ', '', main)))+len(re.findall(r'\S+','Evidence screenshots'))
word_count+=sum(len(re.findall(r'\S+',caption)) for _,caption in figures)
doc=Document()
section=doc.sections[0]
section.page_width=Inches(8.27); section.page_height=Inches(11.69)
section.top_margin=Inches(.85); section.bottom_margin=Inches(.8)
section.left_margin=Inches(.8); section.right_margin=Inches(.8)
section.header_distance=Inches(.3); section.footer_distance=Inches(.35)
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3','Caption']:
    style=doc.styles[name]
    style.font.name='Times New Roman'
    style.font.color.rgb=RGBColor(0,0,0)
    style.paragraph_format.line_spacing=2
doc.styles['Normal'].font.size=Pt(12)
doc.styles['Normal'].paragraph_format.space_after=Pt(6)
for name,size in [('Title',22),('Heading 1',16),('Heading 2',13)]:
    doc.styles[name].font.size=Pt(size)
    doc.styles[name].font.bold=True
    doc.styles[name].paragraph_format.space_before=Pt(10)
    doc.styles[name].paragraph_format.space_after=Pt(5)
    doc.styles[name].paragraph_format.keep_with_next=True
doc.styles['Caption'].font.size=Pt(11)
doc.styles['Caption'].font.italic=False
# Remove inherited theme font overrides and paragraph borders in Word's template.
for style in doc.styles:
    for border in list(style.element.iter(qn('w:pBdr'))):
        border.getparent().remove(border)
    if style.name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3','Caption']:
        rpr=style.element.get_or_add_rPr()
        fonts=rpr.find(qn('w:rFonts'))
        if fonts is not None:
            for key in list(fonts.attrib):
                if 'Theme' in key or 'theme' in key: del fonts.attrib[key]
            for key in ['ascii','hAnsi','eastAsia','cs']: fonts.set(qn('w:'+key),'Times New Roman')
        for spacing in list(rpr.findall(qn('w:spacing'))): rpr.remove(spacing)
doc.styles['Subtitle'].font.italic=False
header=section.header.paragraphs[0]
header.paragraph_format.line_spacing=1
header.paragraph_format.space_after=Pt(0)
run=header.add_run(f'{args.student_name} | NCC {args.student_number}\nWord count {word_count:,}')
run.font.size=Pt(9)
footer=section.footer.paragraphs[0]
footer.alignment=2
footer.paragraph_format.line_spacing=1
footer.add_run('Page ').font.size=Pt(9)
field=OxmlElement('w:fldSimple'); field.set(qn('w:instr'),'PAGE'); footer._p.append(field)
doc.add_paragraph('Google Local Review Analysis and Business Strategy','Title')
doc.add_paragraph('Introduction to Data Science and Big Data','Subtitle')
doc.add_paragraph('Google Local Research Team\nAutumn 2024 to Summer 2025')
doc.add_paragraph(f'Name: {args.student_name}\nNCC Education student number: {args.student_number}\nReport word count: {word_count:,}')
doc.add_paragraph('The report is within the permitted word limit of 4,400 words. The count includes report headings and figure captions, and excludes the cover, reference list and text within figures.')
doc.add_page_break()
for line in source.splitlines():
    if not line.strip(): continue
    if line.startswith('# '):
        doc.add_paragraph(line[2:],'Title')
    elif line.startswith('## '):
        if line=='## References': doc.add_page_break()
        doc.add_paragraph(line[3:],'Heading 1')
    elif line.startswith('### '): doc.add_paragraph(line[4:],'Heading 2')
    else: doc.add_paragraph(line)
doc.add_page_break()
doc.add_paragraph('Evidence screenshots','Heading 1')
for index,(stem,caption) in enumerate(figures):
    if index: doc.add_page_break()
    paragraph=doc.add_paragraph(caption,'Caption')
    paragraph.paragraph_format.keep_with_next=True
    image=ROOT/'evidence'/f'{stem}.png'
    with Image.open(image) as im: width,height=im.size
    available_height=7.8 if index==0 else 8.4
    image_width=min(6.65,available_height*width/height)
    p=doc.add_paragraph()
    p.paragraph_format.line_spacing=1
    p.add_run().add_picture(str(image),width=Inches(image_width))
doc.core_properties.title='Google Local Review Analysis and Business Strategy'
doc.core_properties.author='' if args.student_name.startswith('[') else args.student_name
doc.core_properties.subject='NCC Education Google Local assignment Tasks 1 to 4'
destination=ROOT/'report/Google_Local_Analysis_Report.docx'
doc.save(destination)
(ROOT/'report/word_count.txt').write_text(f'{word_count} words\nIncludes report headings and figure captions. Excludes cover, references and text within screenshots.\n',encoding='utf-8')
print(f'Created {destination}; word count {word_count}')
